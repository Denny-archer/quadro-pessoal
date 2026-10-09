"""Estrutura oficial das planilhas de importação e geração dos arquivos-modelo (CSV e XLSX).

A mesma definição alimenta: a validação, os arquivos-modelo baixados no sistema e a documentação.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import normaliza as N


def _ops(*grupos):
    """[('valor', 'apelido1', 'apelido2'), ...] -> {apelido_normalizado: valor}."""
    out = {}
    for valor, *apelidos in grupos:
        for a in (valor, *apelidos):
            out[N.chave(a)] = valor
    return out


VINCULO = _ops(("efetivo", "empregado efetivo", "concursado", "empregado"),
               ("comissionado", "emprego em comissao", "comissao", "sem vinculo"),
               ("cedido", "cedido de outro orgao", "requisitado"),
               ("estagiario", "estagio", "estagiaria"),
               ("aprendiz", "jovem aprendiz", "menor aprendiz"))
SITUACAO = _ops(("ativo", "ativa", "em exercicio"), ("cedido", "cedida", "cedido a outro orgao"),
                ("inativo", "inativa", "desligado", "desligada"))
REGIME = _ops(("Presencial",), ("Híbrido", "hibrido"), ("Teletrabalho", "home office", "remoto"))
FUNCAO_TIPO = _ops(("comissao", "emprego em comissao", "cargo em comissao", "cc"), ("fg", "funcao gratificada"))
NIVEL = _ops(*[(n, str(i)) for i, n in enumerate(["I", "II", "III", "IV", "V", "VI"], 1)])
CARREIRA = _ops(("analista", "analista / advogado", "advogado"), ("tecnico", "assistente"), ("apoio", "apoio operacional", "auxiliar", "motorista"))
FG_TIPO = _ops(("chefia", "fg chefia"), ("assessor2", "assessor ii", "assessor nivel ii", "fg assessor nivel ii"),
               ("assessor1", "assessor i", "assessor nivel i", "fg assessor nivel i"))
BOLSA = _ops(("medio4", "nivel medio 4h", "medio"), ("tecnico6", "nivel tecnico 6h", "tecnico"),
             ("advocacia6", "profissional advocacia 6h", "advocacia"), ("superior6", "nivel superior 6h", "superior"),
             ("pos6", "pos graduacao 6h", "pos"))
JORNADA = _ops(("4", "4h", "20h"), ("6", "6h", "30h"))
TIPO_LANC = _ops(("ferias", "férias"), ("atestado", "atestado medico", "licenca medica"),
                 ("abono", "folga", "abono anual"), ("viagem", "diaria", "diarias", "viagem a servico"),
                 ("capacitacao", "curso", "treinamento", "evento"), ("banco", "banco de horas", "bh", "banco_horas"))
PERIODO = _ops(("integral", "dia", "do dia", "dia inteiro"), ("matutino", "manha", "m"),
               ("vespertino", "tarde", "v"), ("comparecimento", "declaracao de comparecimento", "comp"))
STATUS_ABONO = _ops(("agendado", "programado"), ("gozado", "usufruido", "utilizado"))


@dataclass
class Coluna:
    nome: str                      # cabeçalho oficial
    descricao: str
    exemplo: str = ""
    obrigatoria: bool = False
    tipo: str = "texto"            # texto|data|hora|numero|inteiro|booleano|enum
    opcoes: dict | None = None
    apelidos: list[str] = field(default_factory=list)
    maximo: int = 200


FUNCIONARIOS = [
    Coluna("matricula", "Matrícula. É a chave: se já existir, o cadastro é atualizado.", "157", apelidos=["mat", "matricula_funcional"], maximo=30),
    Coluna("nome", "Nome completo.", "Maria da Silva Souza", True, apelidos=["nome_completo", "colaborador", "funcionario", "servidor"], maximo=120),
    Coluna("vinculo", "efetivo | comissionado | cedido | estagiario | aprendiz", "efetivo", True, "enum", VINCULO, ["tipo_de_vinculo"]),
    Coluna("situacao", "ativo | cedido | inativo (padrão: ativo)", "ativo", tipo="enum", opcoes=SITUACAO),
    Coluna("admissao", "Data de admissão (dd/mm/aaaa).", "01/03/2016", tipo="data", apelidos=["data_de_admissao", "data_admissao"]),
    Coluna("cargo", "Cargo base (do concurso).", "Assistente Administrativo", apelidos=["cargo_base"], maximo=80),
    Coluna("setor", "Setor / lotação.", "DAGEP", apelidos=["lotacao", "unidade"], maximo=80),
    Coluna("regime", "Presencial | Híbrido | Teletrabalho (padrão: Presencial). A grade semanal do híbrido é definida na tela.", "Presencial", tipo="enum", opcoes=REGIME),
    Coluna("carga_horaria", "Carga horária semanal em horas.", "40", tipo="inteiro", apelidos=["carga", "carga_horaria_semanal", "ch"]),
    Coluna("intervalo_min", "Intervalo de almoço em minutos (padrão 60).", "60", tipo="inteiro", apelidos=["intervalo"]),
    Coluna("entrada", "Horário de entrada presencial (HH:MM).", "08:00", tipo="hora"),
    Coluna("saida", "Horário de saída presencial (HH:MM).", "17:00", tipo="hora"),
    Coluna("entrada_ho", "Entrada no home office (teletrabalho).", "", tipo="hora", apelidos=["entrada_home_office"]),
    Coluna("saida_ho", "Saída do home office (teletrabalho).", "", tipo="hora", apelidos=["saida_home_office"]),
    Coluna("usa_transporte", "Utiliza auxílio-transporte? sim | não", "sim", tipo="booleano", apelidos=["auxilio_transporte", "vale_transporte"]),
    Coluna("valor_transporte", "Valor mensal do auxílio-transporte (R$).", "440,00", tipo="numero", apelidos=["valor_auxilio_transporte"]),
    Coluna("funcao", "Cargo em comissão ou função gratificada.", "Chefe do Setor de Transportes", maximo=160),
    Coluna("funcao_tipo", "comissao | fg", "comissao", tipo="enum", opcoes=FUNCAO_TIPO, apelidos=["tipo_funcao"]),
    Coluna("nivel", "Nível da comissão: I a VI.", "II", tipo="enum", opcoes=NIVEL, apelidos=["nivel_comissao"]),
    Coluna("portaria", "Portaria de nomeação.", "Port. 18/2026", apelidos=["portaria_funcao", "portaria_de_nomeacao"], maximo=60),
    Coluna("ingresso", "Edital / portaria do concurso.", "Edital nº 1/2024", maximo=60),
    Coluna("orgao_origem", "Órgão de origem (cedidos).", "", maximo=80),
    Coluna("cedido_para", "Órgão de destino (situação cedido).", "", maximo=80),
    Coluna("jornada_diaria", "Estágio/aprendizagem: 4 ou 6 horas.", "", tipo="enum", opcoes=JORNADA),
    Coluna("instituicao", "Instituição de ensino (estágio/aprendizagem).", "", apelidos=["instituicao_de_ensino"], maximo=100),
    Coluna("contrato_inicio", "Início do contrato (estágio/aprendizagem, obrigatório para eles).", "", tipo="data", apelidos=["inicio_contrato", "inicio_do_contrato"]),
    Coluna("contrato_fim", "Término do contrato (estágio/aprendizagem, obrigatório para eles).", "", tipo="data", apelidos=["termino_contrato", "fim_contrato", "termino_do_contrato"]),
    Coluna("providencia", "Providência pendente do contrato.", "", maximo=160),
    Coluna("carreira", "analista | tecnico | apoio (vazio = automática pelo cargo).", "", tipo="enum", opcoes=CARREIRA),
    Coluna("nivel_salarial", "Nível salarial 1 a 30 (vazio = pelo tempo de casa).", "", tipo="inteiro"),
    Coluna("fg_tipo", "chefia | assessor2 | assessor1 (vazio = automática).", "", tipo="enum", opcoes=FG_TIPO),
    Coluna("bolsa_tipo", "Estagiário: medio4 | tecnico6 | advocacia6 | superior6 | pos6", "", tipo="enum", opcoes=BOLSA),
]

LANCAMENTOS = [
    Coluna("matricula", "Matrícula do funcionário (preferencial para identificar).", "157", apelidos=["mat"], maximo=30),
    Coluna("nome", "Nome do funcionário (usado se não houver matrícula).", "Maria da Silva Souza", apelidos=["nome_completo", "colaborador", "funcionario", "servidor"], maximo=120),
    Coluna("tipo", "ferias | atestado | abono | viagem | capacitacao | banco_horas", "atestado", True, "enum", TIPO_LANC, ["tipo_de_lancamento", "ocorrencia"]),
    Coluna("inicio", "Data de início (dd/mm/aaaa).", "10/06/2026", True, "data", apelidos=["data", "data_inicio", "data_de_inicio", "de"]),
    Coluna("fim", "Data de término (vazio = mesmo dia do início).", "10/06/2026", tipo="data", apelidos=["termino", "data_fim", "data_de_termino", "ate"]),
    Coluna("periodo", "Só atestado: integral | matutino | vespertino | comparecimento (padrão integral).", "integral", tipo="enum", opcoes=PERIODO, apelidos=["turno"]),
    Coluna("horas", "Banco de horas (obrigatório: + crédito, − débito) ou horas parciais de abono/capacitação.", "", tipo="numero"),
    Coluna("destino", "Viagem/capacitação: destino, evento ou curso.", "", apelidos=["curso", "evento", "local"], maximo=160),
    Coluna("diarias", "Viagem: quantidade de diárias.", "", tipo="numero", apelidos=["qtd_diarias"]),
    Coluna("parcela", "Férias: 1ª, 2ª ou 3ª parcela.", "", maximo=10),
    Coluna("abono_pecuniario", "Férias: dias vendidos (0 a 10).", "", tipo="inteiro", apelidos=["dias_vendidos"]),
    Coluna("aviso_ferias", "Férias: aviso / número do processo.", "", apelidos=["processo"], maximo=120),
    Coluna("referencia_abono", "Abono: ano de referência (ref2024, ref2025, folgas2026...). Vincula ao saldo.", "", maximo=20),
    Coluna("status_abono", "Abono: agendado | gozado (padrão: gozado se já passou).", "", tipo="enum", opcoes=STATUS_ABONO),
    Coluna("obs", "Observação livre. NÃO informe CID ou diagnóstico.", "", apelidos=["observacao", "observacoes"], maximo=600),
]

ESPEC = {"funcionarios": FUNCIONARIOS, "lancamentos": LANCAMENTOS}
EXEMPLOS = {
    "funcionarios": [
        {c.nome: c.exemplo for c in FUNCIONARIOS},
        {"matricula": "E-031", "nome": "João Pereira Lima", "vinculo": "estagiario", "situacao": "ativo", "admissao": "03/02/2026",
         "setor": "DAGEP", "regime": "Presencial", "carga_horaria": "30", "usa_transporte": "não", "jornada_diaria": "6",
         "instituicao": "UnB", "contrato_inicio": "03/02/2026", "contrato_fim": "02/02/2027", "bolsa_tipo": "superior6"},
    ],
    "lancamentos": [
        {"matricula": "157", "nome": "Maria da Silva Souza", "tipo": "ferias", "inicio": "06/07/2026", "fim": "20/07/2026", "parcela": "1ª", "abono_pecuniario": "0"},
        {"matricula": "157", "tipo": "atestado", "inicio": "10/06/2026", "periodo": "vespertino"},
        {"matricula": "49", "tipo": "abono", "inicio": "15/06/2026", "referencia_abono": "ref2025", "status_abono": "gozado"},
        {"matricula": "49", "tipo": "viagem", "inicio": "21/09/2026", "fim": "23/09/2026", "destino": "Fortaleza/CE", "diarias": "2,5"},
        {"matricula": "49", "tipo": "banco_horas", "inicio": "18/09/2026", "horas": "-8", "obs": "folga compensada"},
    ],
}


def mapa_cabecalhos(tipo: str) -> dict[str, Coluna]:
    out = {}
    for c in ESPEC[tipo]:
        for a in (c.nome, *c.apelidos):
            out[N.chave(a)] = c
    return out


def converter(c: Coluna, v):
    if c.tipo == "texto":
        return N.texto(v, c.maximo)
    if c.tipo == "data":
        return N.data(v)
    if c.tipo == "hora":
        return N.hora(v)
    if c.tipo == "numero":
        return N.numero(v)
    if c.tipo == "inteiro":
        return N.inteiro(v)
    if c.tipo == "booleano":
        return N.booleano(v)
    if c.tipo == "enum":
        return N.enum(v, c.opcoes)
    raise ValueError(c.tipo)


def modelo_csv(tipo: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    cols = [c.nome for c in ESPEC[tipo]]
    w.writerow(cols)
    for ex in EXEMPLOS[tipo]:
        w.writerow([ex.get(c, "") for c in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def modelo_xlsx(tipo: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = tipo
    cols = ESPEC[tipo]
    cab = PatternFill("solid", fgColor="122718")
    obr = PatternFill("solid", fgColor="27774A")
    for i, c in enumerate(cols, 1):
        cel = ws.cell(row=1, column=i, value=c.nome)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = obr if c.obrigatoria else cab
        cel.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cel.column_letter].width = max(14, len(c.nome) + 4)
        if c.opcoes:
            valores = sorted(set(c.opcoes.values()))
            if tipo == "lancamentos" and c.nome == "tipo":
                valores = ["ferias", "atestado", "abono", "viagem", "capacitacao", "banco_horas"]
            dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=not c.obrigatoria,
                                showErrorMessage=False)
            dv.add(f"{cel.column_letter}2:{cel.column_letter}2000")
            ws.add_data_validation(dv)
        if c.tipo == "data":
            for r in range(2, 2001):
                ws.cell(row=r, column=i).number_format = "DD/MM/YYYY"
        elif c.tipo in ("texto",):
            for r in range(2, 2001):
                ws.cell(row=r, column=i).number_format = "@"
    for r, ex in enumerate(EXEMPLOS[tipo], 2):
        for i, c in enumerate(cols, 1):
            v = ex.get(c.nome, "")
            if v and c.tipo == "data":
                v = N.data(v)
            ws.cell(row=r, column=i, value=v or None)
    ws.freeze_panes = "A2"
    inst = wb.create_sheet("instrucoes")
    inst.append(["coluna", "obrigatória", "formato / valores aceitos", "exemplo"])
    for c in cols:
        inst.append([c.nome, "sim" if c.obrigatoria else "", c.descricao, c.exemplo])
    for celula in inst[1]:
        celula.font = Font(bold=True)
    inst.column_dimensions["A"].width = 20
    inst.column_dimensions["C"].width = 90
    inst.column_dimensions["D"].width = 28
    inst.append([])
    inst.append(["Apague as linhas de exemplo antes de importar. Cabeçalhos em verde são obrigatórios."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

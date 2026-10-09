"""Lê CSV (UTF-8 ou Windows-1252, separador ; , ou tab) e XLSX, devolvendo as linhas como dicionários."""
from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from . import modelos
from . import normaliza as N


class ArquivoInvalido(ValueError):
    pass


def _decodificar(dados: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ArquivoInvalido("Não foi possível ler o texto do CSV (salve como CSV UTF-8).")


def _tabela_csv(dados: bytes) -> list[list]:
    texto = _decodificar(dados)
    primeira = texto.split("\n", 1)[0]
    sep = max([";", ",", "\t"], key=primeira.count)
    return [l for l in csv.reader(io.StringIO(texto), delimiter=sep)]


def _tabela_xlsx(dados: bytes, tipo: str) -> list[list]:
    try:
        wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except Exception:
        raise ArquivoInvalido("Arquivo XLSX inválido ou protegido por senha.")
    # prefere a aba com o nome do tipo; senão, a primeira que não seja de instruções
    abas = [ws for ws in wb.worksheets if N.chave(ws.title) not in ("instrucoes", "instrucao", "leia_me")]
    ws = next((w for w in abas if N.chave(w.title) == tipo), abas[0] if abas else wb.worksheets[0])
    return [list(r) for r in ws.iter_rows(values_only=True)]


def ler(dados: bytes, nome_arquivo: str, tipo: str):
    """Devolve (colunas reconhecidas, colunas ignoradas, [(nº da linha, {coluna: valor bruto})])."""
    if not dados:
        raise ArquivoInvalido("Arquivo vazio.")
    nome = nome_arquivo.lower()
    if nome.endswith(".xlsx") or nome.endswith(".xlsm"):
        tabela = _tabela_xlsx(dados, tipo)
    elif nome.endswith(".csv") or nome.endswith(".txt"):
        tabela = _tabela_csv(dados)
    elif nome.endswith(".xls"):
        raise ArquivoInvalido("Formato .xls antigo não é aceito: salve como .xlsx ou .csv.")
    else:
        raise ArquivoInvalido("Envie um arquivo .csv ou .xlsx.")

    mapa = modelos.mapa_cabecalhos(tipo)
    # cabeçalho = primeira linha (entre as 10 primeiras) que reconhece ao menos 2 colunas
    idx_cab = next((i for i, l in enumerate(tabela[:10]) if sum(N.chave(c) in mapa for c in l if c is not None) >= 2), None)
    if idx_cab is None:
        esperadas = ", ".join(c.nome for c in modelos.ESPEC[tipo] if c.obrigatoria)
        raise ArquivoInvalido(f"Cabeçalho não encontrado. A primeira linha deve ter os nomes das colunas (obrigatórias: {esperadas}).")
    cab = tabela[idx_cab]
    posicoes, ignoradas, vistas = {}, [], set()
    for i, c in enumerate(cab):
        if c is None or str(c).strip() == "":
            continue
        col = mapa.get(N.chave(c))
        if col is None or col.nome in vistas:
            ignoradas.append(str(c))
        else:
            posicoes[i] = col.nome
            vistas.add(col.nome)
    faltando = [c.nome for c in modelos.ESPEC[tipo] if c.obrigatoria and c.nome not in vistas]
    if tipo == "lancamentos" and not ({"matricula", "nome"} & vistas):
        faltando.append("matricula ou nome")
    if faltando:
        raise ArquivoInvalido(f"Coluna(s) obrigatória(s) ausente(s): {', '.join(faltando)}.")

    linhas = []
    for n, l in enumerate(tabela[idx_cab + 1:], start=idx_cab + 2):
        reg = {nome: (l[i] if i < len(l) else None) for i, nome in posicoes.items()}
        if all(N.vazio(v) for v in reg.values()):
            continue
        linhas.append((n, reg))
    if len(linhas) > 20000:
        raise ArquivoInvalido("Arquivo com mais de 20.000 linhas: divida em partes.")
    return sorted(vistas), ignoradas, linhas
